import base64
import hashlib
import html
import io
import re
from pathlib import Path

from .autopsy_domain import canonical
from .report_brand import design_tokens, report_fonts, report_style

FORMATS = {
    "json": "application/json",
    "html": "text/html",
    "md": "text/markdown",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
}
LABELS = {
    "entry_type": "نوع عنصر الخط الزمني",
    "time_basis": "أساس التوقيت",
    "actor_id": "معرف المسجل",
    "key": "مرجع العنصر",
    "decision": "القرار الموثق",
    "task_id": "المهمة المرتبطة",
    "scope": "نطاق الفحص وحدود الاستنتاج",
    "counts_by_assessment": "تقييم أصول النطاق",
    "earliest_retained_observation": "أقدم ملاحظة محفوظة مدعومة بدليل",
    "counts": "موجز الأدلة المتاحة",
    "normalized_logs": "السجلات المطبعّة",
    "status_at_preparation": "حالة القضية عند إعداد النسخة",
    "created_at": "وقت الإنشاء",
    "closed_at": "وقت الإغلاق",
    "updated_at": "آخر تعديل",
    "priority": "الأولوية",
    "severity": "درجة الخطورة",
    "version": "الإصدار",
    "name": "الاسم",
    "role": "الدور",
    "responsibility": "المسؤولية",
    "user_id": "معرف العضو",
    "created_by": "معرف المؤلف",
    "updated_by": "معرف آخر محرر",
    "added_by": "أضافه العضو",
    "added_at": "وقت حفظ الدليل",
    "id": "المعرف",
    "basis": "أساس حساب القياسات وحدودها",
    "first_evidenced_activity": "أقدم نشاط مثبت بالدليل",
    "detection": "وقت الرصد المثبت",
    "acknowledgement": "وقت إقرار الاستلام",
    "containment": "وقت الاحتواء المثبت",
    "recovery": "وقت التعافي المثبت",
    "detection_delay_seconds": "المدة حتى الرصد بالثواني",
    "acknowledgement_delay_seconds": "المدة حتى الإقرار بالثواني",
    "containment_delay_seconds": "المدة حتى الاحتواء بالثواني",
    "recovery_delay_seconds": "المدة حتى التعافي بالثواني",
    "missing_or_inconsistent_times": "الأوقات الناقصة أو غير المتسقة",
    "evidence_type": "نوع الدليل",
    "source": "المصدر",
    "confidence": "الثقة",
    "details": "تفاصيل التحقيق",
    "legal_hold": "الحجز القانوني",
    "retention_until": "الاحتفاظ حتى",
    "format": "الصيغة",
    "original_filename": "اسم الأصل",
    "parser_version": "إصدار المحلل",
    "accepted": "السجلات المقبولة",
    "rejected": "السجلات المرفوضة",
    "head_sha256": "بصمة رأس سلسلة الحيازة",
    "entries": "عدد حركات الحيازة",
    "summary": "الملخص التنفيذي",
    "case": "بيانات القضية ونطاق التحقيق",
    "context": "الهدف وقرار التنبيه ونتيجة الهجوم",
    "metrics": "قياسات مراحل الاستجابة",
    "members": "المسؤوليات",
    "asset": "الأصول ونطاق الأثر",
    "indicator": "المؤشرات المرصودة والمثبتة",
    "timeline": "الخط الزمني والتعليقات",
    "hypothesis": "الفرضيات والأدلة",
    "foothold": "أقدم نقطة اختراق مثبتة وحدود الاستنتاج",
    "root_cause": "الأسباب الجذرية و5 Whys",
    "playbook": "كتيبات الاستجابة ونتائج الخطوات",
    "response": "الاستجابة والاحتواء والتحقق",
    "escalation": "سجل التصعيد والاستلام والاستجابة",
    "war_room": "قرارات غرفة التحقيق المشتركة",
    "corrective": "التوصيات والإجراءات التصحيحية",
    "reputation": "نتائج السمعة المستوردة يدويًا",
    "analysis": "التحليل المحلي",
    "alerts": "التنبيهات",
    "tasks": "المهام",
    "evidence": "فهرس الأدلة وبصماتها",
    "imports": "الأصول المستوردة وإصدارات المعالجة",
    "custody": "سلامة سلسلة الحيازة",
    "history": "تاريخ حالات القضية",
    "limitations": "حدود التقرير والتحليل",
    "title": "العنوان",
    "reason": "السبب",
    "objective": "الهدف",
    "status": "الحالة",
    "classification": "تصنيف السرية",
    "attack_outcome": "نتيجة الهجوم",
    "alert_verdict": "قرار التنبيه",
    "sha256": "بصمة SHA-256",
    "source_version": "إصدار القضية",
    "generated_at": "وقت إعداد النسخة",
    "evidence_id": "مرجع الدليل",
    "performed_at": "وقت التنفيذ",
    "performed_by": "المنفذ",
    "conclusion": "الاستنتاج",
    "limitations_text": "حدود الاستنتاج",
    "event_time": "وقت الحدث",
    "question": "لماذا؟",
    "answer": "الإجابة",
    "reference": "مرجع القضية",
    "revision": "إصدار التقرير",
    "profile": "سياسة مشاركة التقرير",
    "integrity_valid": "سلامة السلسلة عند إعداد التقرير",
}
ENUM_LABELS = {
    "log": "سجل مصدر",
    "record": "سجل تحقيق أو إجراء",
    "note": "ملاحظة تحقيق",
    "documentation_time": "وقت التوثيق؛ ليس وقت وقوع الهجوم",
    "reported_action_time": "وقت إجراء مسجل بواسطة المنفذ",
    "investigator_observation_time": "وقت ملاحظة أدخله المحقق",
    "hypothesis_time": "توقيت فرضية غير مؤكدة",
    "escalation_time": "وقت إحالة داخل الفريق؛ ليس وقت وقوع الهجوم",
    "source": "وقت مسجل في المصدر",
    "source_year_and_zone_declared": "وقت المصدر مع سنة ومنطقة زمنية معلنتين",
    "source_year_assumed_not_verified": "سنة المصدر مفترضة وغير متحققة",
    "ingestion_fallback_not_observed_time": "وقت استيراد بديل؛ لا يثبت وقت الحدث",
    "new": "مسودة",
    "open": "مفتوحة",
    "investigating": "قيد التحقيق",
    "awaiting_review": "بانتظار المراجعة",
    "closed": "مغلقة",
    "reopened": "أعيد فتحها",
    "successful": "ناجح",
    "blocked": "محجوب أو فاشل",
    "partial": "نجاح جزئي",
    "unknown": "غير معروف",
    "undetermined": "غير محسوم",
    "malicious": "خبيث",
    "benign": "حميد",
    "false_positive": "إيجابي كاذب",
    "internal": "داخلي",
    "confidential": "سري",
    "restricted": "مقيد",
    "critical": "حرج",
    "high": "عالٍ",
    "medium": "متوسط",
    "low": "منخفض",
    "soc_manager": "مدير الفريق",
    "investigator": "محقق",
    "reviewer": "مراجع",
    "analyst": "محلل",
    "owner": "مسؤول القضية",
    "l1": "المستوى الأول",
    "l2": "المستوى الثاني",
    "monitoring": "المراقبة",
}


def display(value) -> str:
    text = str(value) if value is not None else "غير مسجل"
    return re.sub("[\u202a-\u202e\u2066-\u2069]", lambda match: f"\\u{ord(match[0]):04x}", text)


def lines(value, depth=0):
    if isinstance(value, dict):
        for key, item in value.items():
            if isinstance(item, (dict, list)):
                yield depth, "عناصر الخط الزمني" if key == "entries" and isinstance(item, list) else LABELS.get(key, key)
                yield from lines(item, depth + 1)
            else:
                value = "غير قابل للحساب" if item is None and key.endswith("_seconds") else display(item)
                if key in {
                    "status",
                    "status_at_preparation",
                    "severity",
                    "priority",
                    "classification",
                    "attack_outcome",
                    "alert_verdict",
                    "role",
                    "responsibility",
                    "time_basis",
                    "entry_type",
                } and isinstance(item, str):
                    value = ENUM_LABELS.get(item, item)
                yield depth, f"{LABELS.get(key, key)}: {value}"
    elif isinstance(value, list):
        if not value:
            yield depth, "لا توجد عناصر مسجلة"
        for index, item in enumerate(value, 1):
            yield depth, f"— {index} —"
            yield from lines(item, depth + 1)
    else:
        yield depth, display(value)


def sections(payload):
    for key in ("summary", "case", "context", "metrics", "scope", "members", "alerts"):
        if payload.get(key):
            yield LABELS[key], payload[key]
    for kind in (
        "asset",
        "indicator",
        "timeline",
        "hypothesis",
        "foothold",
        "root_cause",
        "playbook",
        "response",
        "escalation",
        "war_room",
        "reputation",
        "analysis",
        "corrective",
    ):
        records = payload.get("records", {}).get(kind, [])
        if records:
            yield LABELS[kind], records
    for title, key in (
        ("الخط الزمني المشترك للأدلة والتحقيق والإجراءات", "chronology"),
        ("الإثراء المحلي ومصادر السمعة وإصداراتها", "observables"),
        ("SLA وجمع الأدلة ومراجعة ما بعد الحادثة", "workflow"),
        ("توثيق MITRE ATT&CK وإصداره وأدلته", "techniques"),
        ("خط زمني للسجلات المختارة كأدلة وتوقيت مصادرها", "evidenced_timeline"),
        ("الملاحظات المشتركة الموروثة", "shared_notes"),
        ("نتائج الإغلاق الموروثة محفوظة بتاريخها", "legacy_closure"),
    ):
        if payload.get(key):
            yield title, payload[key]
    for key in ("tasks", "evidence", "imports", "custody", "history", "limitations"):
        if payload.get(key):
            yield LABELS[key], payload[key]


def line_html(text):
    label, separator, value = text.partition(": ")
    if separator:
        return f'{html.escape(label)}: <bdi dir="auto">{html.escape(value)}</bdi>'
    return f'<bdi dir="auto">{html.escape(text)}</bdi>'


def render_html(payload) -> str:
    case = payload["case"]
    content = []
    outline = []
    for index, (title, value) in enumerate(sections(payload), 1):
        outline.append(f'<li><a href="#report-section-{index}">{html.escape(title)}</a></li>')
        rows = "".join(f'<p class="report-line depth-{min(depth, 3)}">{line_html(text)}</p>' for depth, text in lines(value))
        content.append(f'<section id="report-section-{index}" class="report-section"><h2>{html.escape(title)}</h2>{rows}</section>')
    digest = hashlib.sha256(canonical(payload).encode()).hexdigest()
    title_class = ' class="report-title-long"' if len(case["title"]) > 100 else ""
    cover = f'<header class="report-cover"><p class="report-wordmark" lang="en" dir="ltr">SOCify</p><p dir="ltr">DIGITAL AUTOPSY</p><p>تقرير تحليل وتحقيق في حادثة أمنية</p><h1{title_class}><bdi dir="auto">{html.escape(display(case["title"]))}</bdi></h1><p><bdi dir="ltr">{html.escape(case["reference"])}</bdi> · الإصدار {payload["revision"]}</p></header><p class="report-note">نسخة ثابتة من بيانات القضية عند إعدادها. اعتماد هذه النسخة محفوظ منفصلًا في سجل المنصة. البصمة أدناه تخص JSON المعياري، لا بايتات PDF.</p><p dir="ltr">SHA-256: <bdi>{digest}</bdi></p>'
    navigation = '<nav class="report-outline" aria-label="أقسام التقرير"><h2>محتويات التقرير</h2><ul>' + "".join(outline) + "</ul></nav>"
    style = (
        report_style()
        + ".report-cover{break-inside:avoid}.report-cover h1.report-title-long{font-size:20pt;line-height:1.5}.report-section{break-before:auto}.report-section:first-of-type{break-before:page}.report-line{white-space:pre-wrap;overflow-wrap:anywhere;margin:0 0 10pt;font-size:14pt}.depth-1{padding-inline-start:10pt}.depth-2,.depth-3{padding-inline-start:18pt}"
    )
    return f'<!doctype html><html lang="ar" dir="rtl"><head><meta charset="utf-8"><title>{html.escape(case["reference"])}</title><meta http-equiv="Content-Security-Policy" content="default-src \'none\'; style-src \'unsafe-inline\'; font-src data:; img-src data:"><style>{style}</style></head><body>{cover}{navigation}{"".join(content)}</body></html>'


def bundled_report_asset_fetcher(url, timeout=10, ssl_context=None):
    from weasyprint.urls import URLFetcherResponse, URLFetchingError

    root = Path(__file__).resolve().parent.parent / "static" / "fonts"
    for filename in {
        "Almarai-Regular.woff2",
        "Almarai-Bold.woff2",
        "IBMPlexSans.woff2",
        "BalooBhaijaan2.woff2",
        "IBMPlexSansArabic-Regular.woff2",
        "IBMPlexSansArabic-Bold.woff2",
        *(font["file"] for font in report_fonts()),
    }:
        content = (root / filename).read_bytes()
        if url == "data:font/woff2;base64," + base64.b64encode(content).decode():
            return URLFetcherResponse(url, body=content, headers={"Content-Type": "font/woff2"})
    raise URLFetchingError("Only bundled SOCify report fonts are permitted")


def render_docx(payload) -> bytes:
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Pt(595), Pt(842)
    section.top_margin, section.bottom_margin = Pt(42), Pt(40)
    section.left_margin, section.right_margin = Pt(44), Pt(44)
    normal = document.styles["Normal"]
    tokens = design_tokens()
    ink = tokens["colors"]["paper-text"][0].lstrip("#")
    normal.font.name, normal.font.size, normal.font.color.rgb = tokens["typography"]["en"]["family"], Pt(14), RGBColor.from_string(ink)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.name = "IBM Plex Sans"
        document.styles[name].font.color.rgb = RGBColor.from_string(ink)
        document.styles[name].font.bold = True
        document.styles[name].font.size = Pt({"Title": 28, "Heading 1": 24, "Heading 2": 18}[name])
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        family = tokens["typography"]["ar"]["family"] if name == "Normal" else tokens["typography"]["display"]["ar"]
        document.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), family)
    wordmark = document.add_paragraph().add_run("SOCify")
    wordmark.font.name = tokens["typography"]["display"]["en"]
    wordmark.font.size = Pt(26)
    wordmark.bold = True
    document.add_heading("SOCify Digital Autopsy", 0)
    document.add_heading(display(payload["case"]["title"]), 1)
    document.add_paragraph(f"{payload['case']['reference']} · {payload['revision']}")
    document.add_paragraph("نسخة ثابتة؛ راجع سجل الاعتماد في المنصة. تعرض الأوقات كما سجلت، ولا يعني وقت إعداد التقرير وقت اكتشاف الحادث.")
    for title, value in sections(payload):
        document.add_heading(title, 1)
        for depth, text in lines(value):
            paragraph = document.add_paragraph(text)
            paragraph.paragraph_format.left_indent = Pt(min(depth, 3) * 8)
            bidi = OxmlElement("w:bidi")
            bidi.set(qn("w:val"), "1")
            paragraph._p.get_or_add_pPr().append(bidi)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def render_files(payload):
    from weasyprint import HTML

    html_document = render_html(payload)
    markdown = [
        "# SOCify Digital Autopsy",
        "",
        html.escape(display(payload["case"]["title"])).replace("`", "\\u0060"),
        "",
        "نسخة ثابتة؛ راجع سجل الاعتماد في المنصة.",
        "",
    ]
    for title, value in sections(payload):
        markdown.extend([f"## {title}", "", "```text"])
        markdown.extend(text.replace("`", "\\u0060") for _, text in lines(value))
        markdown.extend(["```", ""])
    return {
        "json": canonical(payload).encode(),
        "html": html_document.encode(),
        "md": "\n".join(markdown).encode(),
        "docx": render_docx(payload),
        "pdf": HTML(string=html_document, url_fetcher=bundled_report_asset_fetcher).write_pdf(),
    }
